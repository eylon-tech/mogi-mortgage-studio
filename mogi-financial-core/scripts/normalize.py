"""Deterministic parsers. No interpolation, inferred publication dates, or missing-to-zero conversion."""
from pathlib import Path
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
import csv, hashlib, io, json, re, sys
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'.python-libs'))
import xlrd
MAP=json.loads((ROOT/'mappings/excel.json').read_text())
VERSION='1.0.0'
MONTHS={s:i+1 for i,s in enumerate(['ינואר','פברואר','מרץ','אפריל','מאי','יוני','יולי','אוגוסט','ספטמבר','אוקטובר','נובמבר','דצמבר'])};MONTHS['מרס']=3

def digest(value): return hashlib.sha256(json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(',',':'),default=str).encode()).hexdigest()
def number(v,error=False):
 if error:return None,'invalid',str(v)
 if v is None or str(v).strip() in ('','..','...','-'):return None,'missing',None if v is None else str(v)
 try:
  n=Decimal(str(v).strip())
  if not n.is_finite():raise InvalidOperation
  return format(n,'f'),'present',str(v)
 except (InvalidOperation,ValueError):return None,'invalid',str(v)
def col(s):return openpyxl.utils.column_index_from_string(s)
def period(s):
 if re.fullmatch(r'\d{4}-\d{2}',s):return date.fromisoformat(s+'-01').isoformat()
 return date.fromisoformat(s).isoformat()
def excel_date(v,mode=0):
 if isinstance(v,(datetime,date)):return v.date() if isinstance(v,datetime) else v
 if isinstance(v,(int,float)) and 20000<v<100000:return xlrd.xldate_as_datetime(v,mode).date()
 if isinstance(v,str):
  try:return datetime.strptime(v.strip(),'%d/%m/%Y').date()
  except ValueError:return None
 return None
class Output:
 def __init__(self,m):
  self.m=m;self.series={};self.rows=[];self.issues=[];self.keys=set();self.publication=None
 def series_id(self,code,label,freq,unit,dims,kind='observed',convention='source_unspecified',mult=0):
  sid=self.m['id']+':'+code
  value=dict(id=sid,source_id=self.m['id'],source_code=code,label=label,frequency=freq,unit=unit,unit_multiplier=mult,value_class=kind,rate_convention=convention,dimensions=dims)
  if sid in self.series and self.series[sid]!=value:raise ValueError('Series metadata drift: '+sid)
  self.series[sid]=value;return sid
 def add(self,sid,p,v,loc,h=-1,effective=None,error=False,release=None):
  key=(sid,p,h,effective)
  if key in self.keys:raise ValueError('Duplicate observation key '+str(key))
  self.keys.add(key)
  value,status,raw=number(v,error)
  self.rows.append(dict(series_id=sid,period=p,horizon_months=h,effective_from=effective,value=value,status=status,raw_value=raw,release_status=release,source_locator=loc))
 def issue(self,code,detail,severity='warning'):self.issues.append(dict(code=code,detail=detail,severity=severity))

def sdmx(o,b):
 rows=csv.DictReader(io.StringIO(b.decode('utf-8-sig')))
 needed={'SERIES_CODE','TIME_PERIOD','OBS_VALUE','FREQ','UNIT_MEASURE','UNIT_MULT'}
 if not needed<=set(rows.fieldnames or []):raise ValueError('SDMX schema changed')
 for n,r in enumerate(rows,2):
  if None in r:raise ValueError('Malformed CSV row')
  dims={k:v for k,v in r.items() if k not in ('TIME_PERIOD','OBS_VALUE','RELEASE_STATUS','COMMENTS')}
  kind='market_model' if o.m['id']=='boi_zcm' else 'observed'
  convention='nominal_annual_percent' if r['SERIES_CODE']=='MNT_RIB_BOI_D' else 'source_unspecified'
  if r['UNIT_MEASURE']!='PT':convention='not_a_rate'
  sid=o.series_id(r['SERIES_CODE'],r['SERIES_CODE'],r['FREQ'],r['UNIT_MEASURE'],dims,kind,convention,int(r['UNIT_MULT']))
  o.add(sid,period(r['TIME_PERIOD']),r['OBS_VALUE'],{'csvRow':n,'column':'OBS_VALUE','comments':r.get('COMMENTS')},release=r.get('RELEASE_STATUS'))
 if o.m['id']=='boi_mortgage':o.issue('AGGREGATE_NOT_CUSTOMER_PRICING',{'reportingEntities':sorted({x['dimensions'].get('REP_ENTITY') for x in o.series.values()}),'note':'Keep SDMX codes. LTV/PTI volume bins are not borrower interest-rate labels.'},'info')

def equator(o,b):
 cfg=MAP['boi_equator'];w=openpyxl.load_workbook(io.BytesIO(b),data_only=True,read_only=True)
 if not {'L01','L02'}<=set(w.sheetnames):raise ValueError('Missing Equator sheets')
 header=cfg['expectedTrackHeaders'];period_sets=[]
 for name in ['L01','L02']:
  s=w[name];data=list(s.iter_rows());c=cfg[name];stride=c['bankStride'];offset=0 if name=='L01' else 1
  banks=[]
  for start in range(3,s.max_column+1,stride):
   bank=data[9][start-1].value
   if not bank:continue
   for k,expected in enumerate(header[offset:]):
    if str(data[11][start+k-1].value).strip()!=expected:raise ValueError(f'Unexpected Equator header {name}:{start+k}')
   banks.append((str(bank).strip(),start))
  if not banks:raise ValueError('No bank columns')
  periods=set()
  for row in data[12:]:
   rawperiod=row[1].value
   if not isinstance(rawperiod,str) or not re.fullmatch(r'\d{4}-\d{2}',rawperiod):continue
   p=period(rawperiod);periods.add(p)
   for bank,start in banks:
    for offset,kind in enumerate(c['measures']):
     cell=row[start+offset-1]
     dims={'bank_name':bank,'track':kind,'measure':'interest' if name=='L01' else 'composition','scope':'bank_monthly_aggregate',**({'rate_definition':'annual_effective_BOI_equator','definition_source':'https://www.boi.org.il/information/bank-paymnts/financial-education/campaigns/boi-equator/mortgage/'} if name=='L01' else {})}
     sid=o.series_id(digest(dims)[:24],bank+' / '+kind+' / '+dims['measure'],'M','percent',dims,convention='effective_annual_percent' if name=='L01' else 'not_a_rate')
     o.add(sid,p,cell.value,{'sheet':name,'cell':cell.coordinate,'periodCell':'B'+str(cell.row)},error=cell.data_type=='e')
  period_sets.append(periods)
 if period_sets[0]!=period_sets[1]:raise ValueError('Equator sheets have different periods')
 d=excel_date(w['L01']['B1'].value)
 if d:o.publication={'date':d.isoformat(),'evidence':'L01!B1 workbook date; not the original publication date of each historical row'}
 o.issue('ROLLING_WINDOW',{'firstAvailable':min(period_sets[0]),'lastAvailable':max(period_sets[0]),'months':len(period_sets[0]),'requestedStart':'2023-03-01','note':'Earlier bank-level observations are absent; aggregate BIR is not a substitute.'})

def forecast(o,b):
 cfg=MAP['boi_forecast451'];w=openpyxl.load_workbook(io.BytesIO(b),data_only=True,read_only=True);keys=[]
 for sheet,c in cfg['sheets'].items():
  s=w[sheet];rows=list(s.iter_rows());horizons=list(range(c['firstHorizon'],c['lastHorizon']+1))
  if [x.value for x in rows[4][3:]]!=horizons:raise ValueError('Forecast horizon header drift '+sheet)
  vintages=set()
  for row in rows[5:]:
   y,month,avg=[x.value for x in row[:3]]
   if not isinstance(y,(int,float)):continue
   if month not in MONTHS or avg not in ('קלנדרי','מדדי'):raise ValueError('Unknown forecast vintage')
   p=f'{int(y):04}-{MONTHS[month]:02}-01';a={'קלנדרי':'calendar','מדדי':'index_period'}[avg];vintages.add((p,a))
   dims={'curve':c['kind'],'averaging':a,'model':'BOI_451_published','scope':'market_expectation'}
   sid=o.series_id(c['kind']+'_'+a,sheet+' / '+avg,'M',c['unit'],dims,'market_model',c['rateConvention'])
   for cell,h in zip(row[3:],horizons):o.add(sid,p,cell.value,{'sheet':sheet,'cell':cell.coordinate,'horizonCell':cell.column_letter+'5'},h,error=cell.data_type=='e')
  keys.append(vintages)
 if not all(k==keys[0] for k in keys):raise ValueError('Forecast sheets do not align by vintage and averaging')
 o.issue('HISTORICAL_RESTATEMENT_POSSIBLE',{'note':'2022 methodology explicitly replaced historical curves. Latest historical workbook is not a point-in-time archive. Original publication times unknown.','methodology':'https://www.boi.org.il/boi_files/Statistics/Estimation%20of%20yields%20from%20government%20bonds.pdf'},'info')

def prepay(o,b):
 w=xlrd.open_workbook(file_contents=b)
 for cfg in MAP[o.m['id']]['tables']:
  s=w.sheet_by_name(cfg['sheet']);pr=col(cfg['periodColumn'])-1;ef=col(cfg['effectiveColumn'])-1
  if len(cfg['valueColumns'])>1:
   if 'ממוצע' not in str(s.cell_value(cfg['headerRow']-1,0)):raise ValueError('Prepayment header drift')
   if 'תוקף' not in str(s.cell_value(cfg['headerRow']-1,ef)):raise ValueError('Missing effective-date header')
  else:
   if 'שיעור' not in str(s.cell_value(cfg['headerRow']-1,2)):raise ValueError('Prepayment legacy header drift')
  headers=[str(s.cell_value(cfg['headerRow']-1,col(c)-1)).strip() for c in cfg['valueColumns']]
  if headers!=cfg['expectedValueHeaders']:raise ValueError('Prepayment tenor headers changed; review mapping')
  for r in range(cfg['firstDataRow']-1,s.nrows):
   d=excel_date(s.cell_value(r,pr),w.datemode);effective=excel_date(s.cell_value(r,ef),w.datemode)
   if not d and not effective:continue
   if not d:raise ValueError(f'Invalid prepayment reporting date {s.name}:{r+1}')
   if not effective:o.issue('MISSING_EFFECTIVE_DATE',{'sheet':s.name,'row':r+1,'reportingDate':d.isoformat(),'note':'Preserved as NULL; not inferred from adjacent rows.'})
   for c,bucket in zip(cfg['valueColumns'],cfg['buckets']):
    cell=s.cell(r,col(c)-1)
    dims={'indexation':'cpi' if o.m['id'].endswith('_linked') else 'unlinked','tenor_bucket_years':bucket,'table_regime':s.name,'scope':'regulatory_reference_not_offer','boundary_regime':'pre_2002_11' if d<date(2002,11,1) else 'from_2002_11','published_bucket_header':str(s.cell_value(cfg['headerRow']-1,col(c)-1)).strip()}
    sid=o.series_id(digest(dims)[:24],s.name+' / '+bucket,'I','percent',dims,convention='effective_annual_percent' if d>=date(2002,11,1) else 'source_unspecified')
    o.add(sid,d.isoformat(),cell.value,{'sheet':s.name,'cell':c+str(r+1),'reportingDate':d.isoformat(),'effectiveCell':cfg['effectiveColumn']+str(r+1)},effective=effective.isoformat() if effective else None,error=cell.ctype==xlrd.XL_CELL_ERROR)
 o.issue('PREPAYMENT_NOT_FEE_CALCULATOR',{'note':'Reference rates and distinct historical tenor regimes only. No automatic statutory fee calculation; pre-November-2002 conventions and bucket boundaries remain distinct; do not auto-convert.'},'info')

def normalize(m):
 b=(ROOT/m['path']).read_bytes()
 if hashlib.sha256(b).hexdigest()!=m['sha256']:raise ValueError('Raw checksum mismatch')
 o=Output(m)
 if m['format']=='cbs_price_snapshot':
  from cbs_parser import parse_cbs
  parse_cbs(o,b)
 elif m['format']=='sdmx_csv':sdmx(o,b)
 elif m['id']=='boi_equator':equator(o,b)
 elif m['id']=='boi_forecast451':forecast(o,b)
 else:prepay(o,b)
 if not o.rows:raise ValueError('Empty dataset')
 out=dict(parserVersion='1.1.0' if m['id']=='boi_equator' else VERSION,source=m,sourceDate=o.publication,series=list(o.series.values()),observations=o.rows,issues=o.issues)
 path=ROOT/'data/normalized'/(m['id']+'.json');path.write_text(json.dumps(out,ensure_ascii=False,separators=(',',':')))
 return dict(source=m['id'],series=len(o.series),observations=len(o.rows),first=min(x['period'] for x in o.rows),last=max(x['period'] for x in o.rows),missing=sum(x['status']=='missing' for x in o.rows),invalid=sum(x['status']=='invalid' for x in o.rows))
if __name__=='__main__':
 import argparse
 args=argparse.ArgumentParser();args.add_argument('--manifest',default='data/manifest.json');args.add_argument('--only');options=args.parse_args()
 report=[]
 for m in json.loads((ROOT/options.manifest).read_text()):
  if options.only and m['id'] not in options.only.split(','):continue
  if m['status']!='downloaded':raise ValueError('Cannot normalize failed download')
  item=normalize(m);report.append(item);print(json.dumps(item))
 (ROOT/'reports/normalization.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
